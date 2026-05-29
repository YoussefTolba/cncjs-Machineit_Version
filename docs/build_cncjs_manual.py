from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_MD = OUT_DIR / "CNCjs_Pendant_User_Manual.md"
OUT_DOCX = OUT_DIR / "CNCjs_Pendant_User_Manual.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(92, 102, 112)
GOLD = RGBColor(122, 90, 0)
RED = RGBColor(155, 28, 28)
BLACK = RGBColor(0, 0, 0)

MANUAL_MD = r"""---
title: CNCjs + Phone Pendant User Manual
machine: Hobby CNC router, MKS DLC V2, GRBL-compatible firmware
units: millimeters
work_coordinate: G54 / P1 by default
updated: 2026-05-17
---

# CNCjs + Phone Pendant User Manual

This manual covers the customized CNCjs workspace and phone pendant used with a GRBL-compatible MKS DLC V2 hobby CNC router. It starts at launching the app and ends with a safe first project workflow.

> **Machine scope:** This guide assumes a wood/plastic/aluminum hobby CNC router, millimeters, a 3D touch probe, and the customized GRBL-only CNCjs build in this workspace.

> **Safety rule:** The software can move a real machine. Keep the router/spindle off while setting up, keep one hand near the emergency stop, and never trust a probing or homing routine until it has been tested in open air.

## Quick Start Route

1. Launch CNCjs with `Launch CNCjs Dev.bat`.
2. Open the desktop workspace at `http://localhost:8080`.
3. Open the phone pendant at `http://<computer-ip>:8000/pendant/`.
4. Connect the MKS DLC V2 serial port with the GRBL controller selected.
5. Unlock and home the machine.
6. Jog to the stock and set the work zero manually or with the Probe widget.
7. Load the first G-code file and inspect the preview.
8. Run an air cut or dry run if the setup is new.
9. Start the job and watch the first pass closely.

## User Startup Guide

Use this section as the normal daily startup procedure. It is written for a user who just walked up to the machine and wants to get CNCjs, the pendant, and the router ready without guessing what comes next.

### 1. Start with the machine safe

1. Make sure the router/spindle is off.
2. Make sure the emergency stop is released but easy to reach.
3. Clear loose tools, screws, wrenches, and offcuts from the table.
4. Check that the tool is tight in the collet.
5. Move the probe cable so it cannot catch on the gantry.

### 2. Power the controller and computer

1. Turn on the CNC controller power.
2. Connect the USB cable between the MKS DLC V2 and the CNCjs computer.
3. Turn on the CNCjs computer if it is not already running.
4. Wait a few seconds for Windows to see the serial port.

### 3. Launch CNCjs

1. Double-click `Launch CNCjs Dev.bat`.
2. Leave the server window open.
3. Open the desktop workspace at `http://localhost:8080`.
4. Note the network IP printed by CNCjs if you want to use the phone pendant.

### 4. Open the phone pendant

1. Connect the phone to the same Wi-Fi network as the CNCjs computer.
2. Open `http://<computer-ip>:8000/pendant/`.
3. Keep the phone screen awake while you are jogging or probing.
4. Use the pendant only while you can see the machine.

### 5. Connect GRBL

1. In the desktop workspace, select the controller serial port.
2. Use the GRBL controller option.
3. Use `115200` baud unless your machine is configured differently.
4. Click Open.
5. Confirm CNCjs shows a real machine state such as Idle or Alarm.

### 6. Unlock and home

1. If GRBL starts in Alarm, inspect the machine before unlocking.
2. Press Unlock only when the machine is physically clear.
3. Press Home Machine if your limits and homing are configured.
4. Use Home X, Home Y, or Home Z only if your firmware supports individual axis homing.

### 7. Jog test before setup

1. Select a small jog step.
2. Jog X a short distance and confirm the machine moves the expected way.
3. Jog Y a short distance and confirm the machine moves the expected way.
4. Jog Z up before any larger X/Y move.
5. Stop and fix the setup if any axis direction feels wrong.

### 8. Prepare the workpiece

1. Clamp the stock firmly.
2. Confirm clamps are outside the expected toolpath.
3. Install the correct tool for the CAM file.
4. Attach the probe if you will use probing.
5. Confirm the probe triggers before sending a probing move.

### 9. Set the work zero

1. Confirm the active WCS is G54 unless you intentionally use another WCS.
2. Use Z probe for the top surface.
3. Use the square probing pad for an edge or corner.
4. Confirm the DRO reads X0, Y0, and Z0 where you expect the job origin.
5. Retract Z to a safe height before loading or starting a job.

### 10. Load and run the first motion carefully

1. Load the G-code file.
2. Inspect the Visualizer before pressing Start.
3. Check that the job extents fit inside the stock.
4. Start with the router/spindle off for an air cut if this setup is new.
5. When you are ready to cut, start the router/spindle and begin the job.
6. Watch the first plunge and first pass closely.

### 11. Shutdown after a job

1. Wait for all motion to stop.
2. Turn off the router/spindle.
3. Raise Z or move to a safe parking position.
4. Close the G-code job if you are finished.
5. Disconnect CNCjs from the controller.
6. Power off the controller after the machine is stationary.

## Before Powering Motion

Use this checklist before any live machine movement:

- Emergency stop is reachable and tested.
- Router/spindle is off while connecting, homing, jogging, and probing.
- Workpiece and wasteboard clamps are clear of the toolpath.
- Probe wire is connected and the probe input changes state when touched.
- Tool is tight in the collet.
- The selected work coordinate system is the one you intend to use, normally G54.
- Jog step is small before moving close to stock.
- The phone running the pendant is on the same network as the CNCjs computer.

## Launch CNCjs

The workspace includes a Windows launcher:

```text
Launch CNCjs Dev.bat
```

When started, the launcher does three things:

1. Builds the CNCjs development files.
2. Starts the desktop web app on port 8080.
3. Starts the CNCjs server on port 8000 and mounts the phone pendant at `/pendant/`.

Open these addresses:

| Screen | Address | Use |
| --- | --- | --- |
| Desktop workspace | `http://localhost:8080` | Main setup, preview, job control, probing, console, and settings |
| Phone pendant | `http://<computer-ip>:8000/pendant/` | Jogging, homing, probing, feed hold, resume, and abort while standing at the machine |

If the phone cannot open the pendant, check that the phone and computer are on the same Wi-Fi network and use the computer IP address printed in the CNCjs server window.

## Connect The Machine

1. Power the MKS DLC V2 controller.
2. Connect the USB cable to the CNCjs computer.
3. In CNCjs, open the connection panel.
4. Select the serial port for the controller.
5. Use GRBL as the controller.
6. Start with `115200` baud unless the controller was configured differently.
7. Click Open and wait for CNCjs to show a machine state.

This customized build is intended for GRBL-compatible firmware. Laser, Marlin, Smoothie, and TinyG paths were removed from the normal machine workflow.

Common states:

| State | Meaning | Normal action |
| --- | --- | --- |
| Idle | GRBL is connected and ready | Safe to jog, probe, or start a job |
| Alarm | GRBL needs attention, often after power-up or a limit event | Inspect the machine, then unlock or home |
| Run | A program or motion command is active | Watch the machine |
| Hold | Motion is paused | Resume or abort after checking the machine |
| Disconnected | CNCjs is not connected to GRBL | Check USB, port, power, and baud rate |

## Workspace Overview

The normal desktop/tablet workspace is the best place to set up a job. Use it when you need a larger preview, console output, or detailed controls.

| Area | What it is for |
| --- | --- |
| Axes / DRO | Read machine position and work position, jog, home, and zero axes |
| Visualizer | Inspect the loaded G-code toolpath before running |
| Probe widget | Z probing, X/Y edge probing, and corner probing with the square pad |
| Tool tab | Controls M6 tool-change behavior and optional tool probing positions |
| Console | Send or inspect raw GRBL commands and responses |
| Jobs / Files | Load, preview, start, pause, and stop G-code jobs |
| Settings | App settings, workspace records, events, and the GRBL Settings wizard |

## GRBL Settings Wizard

Use the GRBL Settings wizard when you need to inspect, back up, edit, or restore firmware settings. Open it from the desktop workspace at Settings > GRBL Settings after connecting to a GRBL controller.

The wizard reads settings with `$$` and writes changes as normal GRBL `$number=value` commands. Do not use it during a running job.

| Tab | Use |
| --- | --- |
| Backup | Save the current settings to a backup file, load a previous backup file, and flash settings from that file |
| Motion | Edit travel resolution, maximum rates, acceleration, travel limits, and steps/mm calibration |
| Homing & Limits | Edit limit switch, homing, debounce, search, feed, and pull-off settings |
| Spindle & Laser | Edit spindle speed range and laser mode |
| Advanced | Edit pulse, invert, report, junction, arc, and unit-reporting settings |

### Back Up GRBL Settings

1. Connect to the GRBL controller.
2. Open Settings > GRBL Settings.
3. Click Refresh if the settings list is empty or stale.
4. Open the Backup tab.
5. Click Back Up.
6. Keep the downloaded `grbl-settings-YYYYMMDD-HHMMSS.txt` file somewhere safe.

The backup file contains readable comments plus flashable `$number=value` lines. The wizard also displays the generated or loaded file so you can inspect it before using it.

### Restore A Backup File

1. Connect to the same GRBL controller or a compatible replacement controller.
2. Open Settings > GRBL Settings > Backup.
3. Click Choose Backup File and select the saved settings file.
4. Confirm the displayed file contains the settings you expect.
5. Click Flash Backup.
6. Confirm the inline warning.
7. Wait for the wizard to send the settings and read `$$` back from GRBL.

Only flash backup files you trust. A GRBL settings file can change steps/mm, axis direction, limit behavior, homing behavior, spindle speed mapping, and laser mode.

### Safe Editing Rules

1. Back up the current settings before making changes.
2. Change one group of settings at a time.
3. Keep the machine idle and the router/spindle off while changing firmware settings.
4. Test jog direction, travel distance, homing, limits, and probing before cutting again.
5. If a change makes the machine behave incorrectly, reload the backup file and flash it back.

## Phone Pendant Overview

The phone pendant is for standing near the CNC while you jog, touch off, or pause motion. It mirrors the important machine actions without the dense desktop UI.

Use the pendant for:

- Jogging X, Y, and Z.
- Homing all axes or one axis at a time.
- Running the 9-button probing pad.
- Changing probing variables.
- Feed hold, resume, and abort.

Do not use the pendant as a substitute for watching the machine. Keep the screen awake during motion and do not put the phone in a pocket while connected.

## Homing

Use Home Machine after power-up if your GRBL configuration supports homing. Homing establishes machine coordinates, which makes safe machine-coordinate moves possible.

The individual axis buttons use these command sequences:

| Button | Commands |
| --- | --- |
| Home X | `G53 G0 Z-5`, `%wait`, `$HX` |
| Home Y | `G53 G0 Z-5`, `%wait`, `$HY` |
| Home Z | `$HZ` |

> **Important:** `G53 G0 Z-5` moves in machine coordinates before X or Y homing. Use Home X or Home Y only if machine-coordinate Z-5 is a safe raised position on your router.

If your GRBL firmware does not support `$HX`, `$HY`, or `$HZ`, use the normal Home Machine command or update the firmware homing setup before relying on individual axis homing.

## Jogging

Jogging moves the machine manually. Use small steps near the workpiece and larger steps only when the tool is clear.

Good jogging habits:

- Home first when possible.
- Raise Z before large X/Y moves.
- Start with a small jog step near clamps, stock, or the probe.
- Watch the actual machine, not only the DRO.
- If an axis moves opposite of what you expect, stop and fix the machine direction before cutting.

On the phone pendant, the jog pad is arranged to match the master workspace. X/Y moves are grouped as a square pad and Z moves sit beside it.

## Probe Widget Basics

The Probe widget has a Z button and a square 3x3 stock pad. The pad is visual so you can select the edge or corner you are touching.

Square map:

| Back Left | Back Edge `Y-` | Back Right |
| --- | --- | --- |
| Left Edge `X+` | Z | Right Edge `X-` |
| Front Left | Front Edge `Y+` | Front Right |

Corner direction map:

| Corner | X probe | Y probe |
| --- | --- | --- |
| Front Left | X+ | Y+ |
| Front Right | X- | Y+ |
| Back Left | X+ | Y- |
| Back Right | X- | Y- |

All entered distances are positive numbers. The selected button decides whether the routine sends a positive or negative axis move.

## Probe Variables

The workspace and pendant use the same visible probing variables:

| Variable | Meaning | Typical starting value |
| --- | --- | --- |
| Probe Speed | Feed used by `G38.2`, in mm/min | 20 |
| Probe Depth | Search distance for the probing move | 10 |
| Z Probe Offset | Z value applied after a Z touch | 0 for direct 3D probe touch |
| XY Probe Offset | Probe radius used when touching an X/Y edge | Probe ball radius |
| Retraction Distance | Move away from the touched surface after probing | 4 or 5 |

Use conservative values first. A smaller probe depth reduces crash risk if the probe starts in the wrong place.

## Basic Probing Cycles

The simple working cycles are intentionally direct. Each one probes, applies the active G54/P1 work offset, then retracts away from the touched surface.

Z probe:

```gcode
G91
G38.2 Z-{probeDepth} F{probeSpeed}
G90
G10 L20 P1 Z{zProbeOffset}
G91
G0 Z{retractionDistance}
G90
```

X+ probe:

```gcode
G91
G38.2 X{probeDepth} F{probeSpeed}
G90
G10 L20 P1 X-{xyProbeOffset}
G91
G0 X-{retractionDistance}
G90
```

X- probe:

```gcode
G91
G38.2 X-{probeDepth} F{probeSpeed}
G90
G10 L20 P1 X{xyProbeOffset}
G91
G0 X{retractionDistance}
G90
```

Y+ probe:

```gcode
G91
G38.2 Y{probeDepth} F{probeSpeed}
G90
G10 L20 P1 Y-{xyProbeOffset}
G91
G0 Y-{retractionDistance}
G90
```

Y- probe:

```gcode
G91
G38.2 Y-{probeDepth} F{probeSpeed}
G90
G10 L20 P1 Y{xyProbeOffset}
G91
G0 Y{retractionDistance}
G90
```

If a probe fails to trigger, do not continue the setup. Stop, retract safely, inspect the wiring and starting position, then try again.

## Corner Probing

Corner probing starts with a Z touch-off, retracts, moves away from the selected X edge, lowers to the side-probing height, probes X, moves to a Y starting position, probes Y, raises to Z+5, then returns to X0 Y0.

The first corner that was built and tested was Front Left:

1. Probe Z normally.
2. Set Z0 using the Z probe offset.
3. Retract Z by the retraction distance.
4. Move in negative X by the probe depth value.
5. Move to the side probing Z height.
6. Probe X+ and set X0 using negative XY probe offset.
7. Retract from the X wall.
8. Raise to the Z probe retract height.
9. Move to X0+10.
10. Move in negative Y by the probe depth value.
11. Move to the side probing Z height.
12. Probe Y+ and set Y0 using negative XY probe offset.
13. Retract from the Y wall.
14. Move to Z+5.
15. Move to X0 Y0.

The other corners use the same shape with changed X/Y directions:

| Corner | Backoff before X probe | X probe | Y setup move | Y probe |
| --- | --- | --- | --- | --- |
| Front Left | X negative | X+ | Y negative | Y+ |
| Front Right | X positive | X- | Y negative | Y+ |
| Back Left | X negative | X+ | Y positive | Y- |
| Back Right | X positive | X- | Y positive | Y- |

Before using a corner routine, jog the probe so the routine has room to search in both directions. The routine cannot know where clamps, screws, or stock overhangs are.

## Tool Tab

The Tool tab controls what CNCjs does when loaded G-code contains an M6 tool-change command.

Common policies:

| Policy | What happens |
| --- | --- |
| Ignore M6 | CNCjs ignores tool-change commands. This is the safe default for simple hobby jobs. |
| Send M6 | CNCjs sends M6 to the controller. Use only if your firmware handles it correctly. |
| Manual Tool Change (WCS) | CNCjs pauses, moves to the configured tool-change/probe positions, and adjusts work Z through the WCS method. |
| Manual Tool Change (TLO) | CNCjs uses a tool length offset style workflow. |
| Manual Tool Change (Custom Probing) | CNCjs uses the custom probing command configured in the Tool tab. |

Use the Tool tab only after the tool-change position and tool-probe position have been tested in machine coordinates. A bad machine-coordinate tool-change position can move through clamps or stock.

## First Project Walkthrough

This is a conservative first-cut workflow for a simple CAM file.

### 1. Prepare the CAM file

- Set CAM units to millimeters.
- Use a GRBL-compatible post processor.
- Put the work origin where you can probe it. A common choice is the front-left top of the stock.
- Keep the first project shallow and simple.
- Save the G-code file somewhere easy to find.

### 2. Prepare the machine

1. Mount the workpiece securely.
2. Install the correct tool.
3. Keep the router/spindle off.
4. Connect the probe if you will probe.
5. Move the dust hose and clamps clear of the expected toolpath.

### 3. Launch and connect

1. Run `Launch CNCjs Dev.bat`.
2. Open the desktop workspace.
3. Connect the serial port.
4. Unlock or home if GRBL starts in Alarm.
5. Open the phone pendant if you will stand at the machine.

### 4. Home and jog

1. Press Home Machine if the machine supports full homing.
2. Jog near the stock with large steps while Z is high.
3. Switch to small jog steps near the final setup position.
4. Stop if an axis moves in the wrong direction.

### 5. Set work zero

For a front-left top origin:

1. Jog the probe or tool near the top surface.
2. Run Z probe.
3. Jog to the selected front-left corner start area.
4. Run Front Left corner probe, or use X+ and Y+ edge probes manually.
5. Confirm the DRO reads close to X0, Y0, Z0 when the tool is at the work origin.

If you are not using the probe, jog manually and use the axis zero buttons instead.

### 6. Load and inspect G-code

1. Load the G-code file in CNCjs.
2. Inspect the Visualizer.
3. Confirm the toolpath is above the stock where expected.
4. Confirm the job extents fit inside the workpiece.
5. Look for moves that go through clamps.

### 7. Run the first job

1. Put on eye and ear protection.
2. Start dust collection if used.
3. Start the router/spindle if your workflow requires manual spindle control.
4. Press Start.
5. Watch the first plunge and first pass closely.
6. Use Feed Hold if anything looks wrong but there is no immediate danger.
7. Use Abort or the physical emergency stop if there is immediate danger.

### 8. Finish safely

1. Wait until motion stops.
2. Turn off the router/spindle.
3. Raise Z if needed.
4. Remove dust and chips before unclamping.
5. Save any notes about feed, speed, zeroing, or probing values that worked.

## Troubleshooting

| Problem | Likely cause | What to do |
| --- | --- | --- |
| Alarm after probing | Probe stayed triggered, retract was too small, limit was hit, or GRBL reported a probing error | Stop, retract manually only when safe, clear the probe, unlock, and reduce probe depth or speed |
| Probe does not stop | Probe wiring or GRBL probe input is wrong | Test the probe input before motion and verify `$6` invert setting if needed |
| Probe stops too late | Feed too high, controller latency, loose probe, or mechanical flex | Use slow probe speed and short depth, check probe mount and wiring |
| Corner result is wrong | Wrong corner selected or machine front/back convention misunderstood | Use the visual square map and verify with air probing |
| Pendant will not load | Wrong IP, phone not on same Wi-Fi, server not running, or firewall blocked port 8000 | Use the IP printed by CNCjs, keep the server window open, and allow local network access |
| Serial port missing | USB cable, driver, controller power, or port already open elsewhere | Replug USB, power cycle controller, close other sender software |
| `$HX` or `$HY` fails | Firmware does not support individual axis homing | Use Home Machine or update GRBL firmware/configuration |
| Job starts in wrong place | Work zero was set on the wrong corner or wrong WCS | Stop, verify G54, re-zero, reload/preview the job |

## Good Defaults To Start With

| Setting | Conservative value |
| --- | --- |
| Probe Speed | 20 mm/min |
| Probe Depth | 10 mm |
| Z Probe Offset | 0 mm for direct stock touch |
| XY Probe Offset | Actual probe ball radius |
| Retraction Distance | 4 mm or 5 mm |
| Work Coordinate | G54 / P1 |

Change one value at a time and write down what changed. That makes it much easier to find the cause if a probe or first cut goes wrong.

## Operator Notes

- Keep the normal workspace open even when using the pendant. It gives you the larger visualizer and console.
- Use the pendant for close-up control, not for unattended operation.
- Always test a new probing routine in free space first.
- When in doubt, raise Z, feed hold, and think before sending the next command.
- For a first project, choose a file that removes very little material and has plenty of clearance from clamps.
"""


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color="D7DBE2", size="8", space="4"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def shade_element(element, fill):
    tc_pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else None
    if tc_pr is None:
        p_pr = element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        p_pr.append(shd)
        return
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_paragraph_format(paragraph, before=0, after=6, line=1.25, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_together = True


def create_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 24, NAVY, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(9.5)
    code.font.color.rgb = RGBColor(28, 35, 43)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.1)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.1
    code.paragraph_format.keep_together = True


def add_inline_markdown(paragraph, text, base_size=11, base_color=BLACK):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", base_size - 0.5, RGBColor(28, 35, 43))
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, "Calibri", base_size, base_color, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, "Calibri", base_size, base_color)


def add_callout(doc, lines):
    text = " ".join(line.lstrip("> ").strip() for line in lines)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_fill(cell, "FFF8E8")
    p = cell.paragraphs[0]
    set_paragraph_format(p, after=0, line=1.25, keep=True)
    add_inline_markdown(p, text)
    doc.add_paragraph()


def add_markdown_table(doc, table_lines):
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return

    col_count = max(len(r) for r in rows)
    widths = [int(9360 / col_count)] * col_count
    widths[-1] += 9360 - sum(widths)

    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    set_table_width(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_format(p, after=0, line=1.15, keep=True)
            add_inline_markdown(p, row[c_idx] if c_idx < len(row) else "", base_size=10.5)
            if r_idx == 0:
                set_cell_fill(cell, "E8EEF5")
                for run in p.runs:
                    run.bold = True
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, after=4)


def add_code_block(doc, code_lines):
    for i, line in enumerate(code_lines):
        p = doc.add_paragraph(style="Code Block")
        if i == 0:
            p.paragraph_format.space_before = Pt(4)
        if i == len(code_lines) - 1:
            p.paragraph_format.space_after = Pt(8)
        shade_element(p._p, "F4F6F9")
        run = p.add_run(line if line else " ")
        set_run_font(run, "Consolas", 9.5, RGBColor(28, 35, 43))


def parse_body_markdown(md):
    lines = md.splitlines()
    body = []
    in_front_matter = False
    consumed_title = False
    for i, line in enumerate(lines):
        if i == 0 and line.strip() == "---":
            in_front_matter = True
            continue
        if in_front_matter:
            if line.strip() == "---":
                in_front_matter = False
            continue
        if not consumed_title and line.startswith("# "):
            consumed_title = True
            continue
        body.append(line)
    return body


def render_markdown_to_docx(doc, md):
    lines = parse_body_markdown(md)
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue

        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i])
                i += 1
            add_callout(doc, block)
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            p = doc.add_paragraph(style=f"Heading {level}" if level <= 3 else "Heading 3")
            add_inline_markdown(p, heading.group(2), base_size={1: 16, 2: 13, 3: 12}.get(level, 12), base_color=BLUE if level < 3 else DARK_BLUE)
            i += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, numbered.group(1))
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("- ")
                or re.match(r"^\d+\.\s+", nxt)
                or nxt.startswith("|")
                or nxt.startswith(">")
                or nxt.startswith("```")
            ):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        set_paragraph_format(p, keep=True)
        add_inline_markdown(p, " ".join(paragraph_lines))


def add_cover(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.text = "CNCjs + Phone Pendant"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], "Calibri", 9, MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = "GRBL hobby router operator guide"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.runs[0], "Calibri", 9, MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(64)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Operator Manual")
    set_run_font(run, "Calibri", 12, GOLD, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("CNCjs + Phone Pendant User Manual")
    set_run_font(run, "Calibri", 26, NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("From launching the app to running your first GRBL router project")
    set_run_font(run, "Calibri", 13, MUTED)

    rule = doc.add_paragraph()
    paragraph_border_bottom(rule, color="D7DBE2", size="12", space="4")
    rule.paragraph_format.space_after = Pt(18)

    metadata = [
        ("Machine", "Hobby CNC router with MKS DLC V2 / GRBL-compatible firmware"),
        ("Units", "Millimeters"),
        ("Default WCS", "G54 / P1"),
        ("Probe", "3D touch probe"),
        ("Updated", "May 17, 2026"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_width(table, [2700, 6660])
    for idx, (label, value) in enumerate(metadata):
        label_cell = table.cell(idx, 0)
        value_cell = table.cell(idx, 1)
        set_cell_fill(label_cell, "E8EEF5")
        for cell, text, bold in ((label_cell, label, True), (value_cell, value, False)):
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_format(p, after=0, line=1.15, keep=True)
            run = p.add_run(text)
            set_run_font(run, "Calibri", 10.5, BLACK, bold=bold)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Keep this manual near the machine. Test every motion routine in open air before cutting material.")
    set_run_font(run, "Calibri", 10.5, RED, bold=True)

    doc.add_page_break()


def build_docx():
    doc = Document()
    create_styles(doc)
    add_cover(doc)
    render_markdown_to_docx(doc, MANUAL_MD)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

    doc.save(OUT_DOCX)


def main():
    OUT_DIR.mkdir(exist_ok=True)
    OUT_MD.write_text(MANUAL_MD, encoding="utf-8", newline="\n")
    build_docx()
    print(OUT_MD)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
